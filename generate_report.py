"""Generate the Week 1 beauty and wellness analysis proposal as a Word report."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError as error:
    raise SystemExit(
        "python-docx is required. Activate the project virtual environment and "
        "run: pip install -r requirements.txt"
    ) from error

from src.report_utils import add_hyperlink, add_page_field, add_toc_field
from src.research_utils import read_csv_records, validate_source_records


PROJECT_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = PROJECT_DIR / "docs" / "week1_beauty_wellness_problem_definition.docx"
SOURCE_LOG = PROJECT_DIR / "docs" / "research_sources.csv"
VARIABLE_DICTIONARY = PROJECT_DIR / "docs" / "variable_dictionary.csv"

NAVY = "17324D"
TEAL = "138A8A"
PALE_TEAL = "E8F4F3"
MID_GREY = "596775"
LIGHT_GREY = "F3F5F7"
WHITE = "FFFFFF"
TABLE_COUNTER = 0


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def add_bullet(document: Document, text: str, level: int = 0) -> None:
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def add_numbered(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def add_callout(document: Document, label: str, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, PALE_TEAL)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    lead = paragraph.add_run(f"{label}  ")
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(TEAL)
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(document: Document, headers: list[str], rows: list[list[str]], font_size: float = 8.5) -> None:
    global TABLE_COUNTER
    TABLE_COUNTER += 1
    caption = document.add_paragraph(style="Caption")
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(3)
    caption.add_run(f"Table {TABLE_COUNTER}. {' / '.join(headers)}")
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, heading in enumerate(headers):
        cell = header_row.cells[index]
        shade_cell(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(heading)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            cell = cells[column_index]
            if row_index % 2 == 1:
                shade_cell(cell, LIGHT_GREY)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            run.font.size = Pt(font_size)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_heading(document: Document, title: str, level: int = 1) -> None:
    paragraph = document.add_heading(title, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_body(document: Document, text: str, lead: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    if lead and text.startswith(lead):
        paragraph.add_run(lead).bold = True
        paragraph.add_run(text[len(lead) :])
    else:
        paragraph.add_run(text)


def set_document_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("263746")
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size, color in (
        ("Title", 31, NAVY),
        ("Subtitle", 15, TEAL),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 12, TEAL),
        ("Heading 3", 10, NAVY),
        ("Caption", 8, MID_GREY),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name in ("Title", "Heading 1") else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if style_name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(10 if style_name == "Heading 1" else 6)
            style.paragraph_format.space_after = Pt(4)


def configure_section(section) -> None:
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("WEEK 1 PROJECT  |  BEAUTY + WELLNESS ANALYTICS")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    label = footer.add_run("PROPOSED ANALYSIS  •  ")
    label.font.size = Pt(8)
    label.font.color.rgb = RGBColor.from_string(MID_GREY)
    add_page_field(footer)


def set_update_fields(document: Document) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def validate_source_log() -> list[dict[str, str]]:
    required = {f"S{number}" for number in range(1, 8)} | {"S11", "S12"}
    rows = read_csv_records(SOURCE_LOG)
    validate_source_records(
        rows,
        required_ids=required,
        dated_ids={f"S{number}" for number in range(1, 8)},
    )
    return rows


def load_variable_dictionary() -> list[dict[str, str]]:
    rows = read_csv_records(VARIABLE_DICTIONARY)
    required_columns = {
        "variable",
        "definition",
        "data_type_example",
        "role",
        "source",
        "required_optional",
        "quality_checks",
        "privacy_notes",
    }
    if not rows or not required_columns.issubset(rows[0]):
        raise ValueError("Variable dictionary is empty or missing required columns.")
    return rows


def add_cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(66)
    run = paragraph.add_run("WEEK 1  |  INTERNSHIP PROJECT")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(12)
    title.add_run("Beauty & Wellness\nCustomer Purchase and Retention Analysis")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("Evidence-led research proposal and analysis blueprint")
    document.add_paragraph()

    add_callout(
        document,
        "PROJECT STATUS",
        "Proposed analysis only. No customer dataset has been analyzed and no "
        "hypothesis has been tested. This report contains no customer findings.",
    )
    details = document.add_paragraph()
    details.paragraph_format.space_before = Pt(24)
    details.paragraph_format.line_spacing = 1.35
    details.add_run("Prepared: October 10, 2026\n")
    details.add_run("Planned effort: 34 hours (within the 30–35-hour brief)\n")
    details.add_run("Primary output: Customer-respectful, measurable retention analysis plan")
    footer = document.add_paragraph()
    footer.paragraph_format.space_before = Pt(36)
    run = footer.add_run(
        "Industry context is attributed to public McKinsey publications; it is not "
        "evidence about a specific company or its customers."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)
    document.add_page_break()


def add_table_of_contents(document: Document) -> None:
    add_heading(document, "Contents", level=1)
    add_body(
        document,
        "The Word table of contents is field-based. If it is not populated when "
        "opened, right-click it and select Update Field, then Update entire table.",
    )
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    add_toc_field(paragraph)
    document.add_page_break()


def add_main_content(document: Document) -> None:
    add_heading(document, "1. Executive summary")
    add_body(
        document,
        "This Week 1 project proposes a disciplined way to understand customer "
        "purchase and repeat behavior in beauty and wellness. It begins with "
        "clear business definitions, approved transaction data, and cohort-aware "
        "descriptive analysis. Modeling is optional and conditional on data "
        "quality, sufficient follow-up, governance approval, and a decision it "
        "can improve. No customer findings are reported because no customer data "
        "were supplied.",
    )
    add_body(
        document,
        "The intended decision is not simply who to target. It is whether "
        "validated patterns can support more relevant service or communication, "
        "measured against incremental customer and business outcomes while "
        "protecting privacy, consent, and customer experience.",
    )
    add_table(
        document,
        ["Project element", "Week 1 proposal"],
        [
            ["Business question", "Can governed purchase data clarify repeat behavior and inform testable, customer-respectful retention actions?"],
            ["Analysis sequence", "Governance and data checks → descriptive EDA → comparable cohorts → optional time-aware modeling → controlled tests."],
            ["Customer outcomes", "Repeat purchase and time to second order, with explicit eligibility and complete follow-up."],
            ["Evidence boundary", "Published industry context is not company evidence; proposed hypotheses remain untested."],
            ["Planned effort", "34 hours, with the modeling block reallocated if its readiness gate fails."],
        ],
    )
    add_callout(
        document,
        "READ THIS FIRST",
        "All proposed windows, variables, methods, and interventions require sponsor "
        "and data-owner approval. Do not present them as measured results.",
    )

    add_heading(document, "2. Industry context and business relevance")
    add_body(
        document,
        "The selected topic is relevant because beauty and wellness span routine "
        "purchases, replenishment, discovery, and service experiences. McKinsey's "
        "2026 beauty article describes shifting categories, channels, and shopping "
        "behavior and states that it expects the global beauty market to grow by "
        "5 percent annually through 2030 [1]. Its 2025 beauty report also discusses "
        "value scrutiny and projects online channels to approach one-third of "
        "global beauty sales by 2030, from 26 percent in 2024 [2]. These are "
        "McKinsey's attributed market outlooks, not forecasts produced by this project.",
    )
    add_body(
        document,
        "McKinsey's 2025 Future of Wellness article describes a survey of more "
        "than 9,000 consumers across China, Germany, the United Kingdom, and the "
        "United States, covering six dimensions: health, sleep, nutrition, "
        "fitness, appearance, and mindfulness [3]. This illustrates the breadth "
        "of the consumer research context; it does not measure the customers of "
        "a specific business. Market-level context motivates investigation, but "
        "retention opportunity must be established in the approved business data.",
    )
    add_callout(
        document,
        "BUSINESS INTERPRETATION",
        "Market growth does not guarantee customer retention, profitable repeat "
        "purchases, or campaign lift for an individual organization.",
    )
    add_heading(document, "2.1 Academic research and methodological rationale", level=2)
    add_body(
        document,
        "The industry reports establish context; academic journal and "
        "customer-analytics research informs how this proposal defines and "
        "tests customer behavior. These studies do not provide results for the "
        "business under study.",
    )
    add_table(
        document,
        ["Research", "Contribution to this proposal", "Transfer boundary"],
        [
            [
                "Perret & Schwientek (2025), Digital [4]",
                "Beauty-specific online-survey research on AR/AI cosmetics connects experience and satisfaction concepts with purchase, advocacy, and repurchase intentions.",
                "Narrow beauty-tech context; intentions are not observed repeat transactions or causal intervention effects. Review full methods before relying on detailed findings.",
            ],
            [
                "Fader, Hardie & Lee (2005), Journal of Marketing Research [5]",
                "Formal RFM/customer lifetime value model with holdout tests; motivates RFM candidate features and out-of-sample validation.",
                "Applied to online music customers; assumptions must be checked in beauty/wellness data.",
            ],
            [
                "Fader, Hardie & Shang (2010), Marketing Science [6]",
                "Discrete-time model of noncontractual repeat behavior and latent dropout; motivates distinguishing observed inactivity from confirmed churn.",
                "Applied to nonprofit donations; not a beauty retail validation, and no such model is fitted here.",
            ],
            [
                "Gupta, Lehmann & Stuart (2004), Journal of Marketing Research [7]",
                "Frames customer value using expected discounted future earnings; motivates linking retention outcomes to actual margin and cost.",
                "Illustrative firm-valuation evidence is context-specific; no reported percentage effect is adopted as a target or forecast.",
            ],
        ],
        font_size=7.8,
    )
    add_body(
        document,
        "Research synthesis: use market evidence to frame the question, "
        "beauty-specific survey research to motivate customer-experience "
        "constructs, and marketing-science methods to guide behavioral measures "
        "and validation. The primary outcomes remain approved customer "
        "transactions, not survey intention or external model estimates. "
        "Bibliographic records, abstract-level verification, and transfer "
        "limits are documented in the source log.",
    )
    document.add_page_break()

    add_heading(document, "3. Business problem, scope, and objectives")
    add_body(
        document,
        "Business problem: a beauty and wellness retailer lacks a systematic, "
        "evidence-based way to identify customer purchasing patterns and "
        "understand which observable factors may be associated with repeat "
        "purchases and retention. The decision challenge is how to improve "
        "understanding of customer repeat behavior to support more effective "
        "marketing choices, without assuming a retention problem has already "
        "been demonstrated.",
    )
    add_body(
        document,
        "Data science problem: given approved customer transaction and "
        "engagement data, how can purchase patterns be described and which "
        "available customer or transaction characteristics are associated with "
        "a predefined future repeat-purchase outcome? Unit of analysis: one "
        "pseudonymous customer at an agreed prediction/measurement cutoff; "
        "transaction lines are first aggregated to distinct qualifying orders. "
        "Where no usable customer key exists, the question must be narrowed to "
        "order-level descriptive analysis rather than pretending repeat "
        "behavior is observable.",
    )
    add_heading(document, "3.1 Objectives", level=2)
    for item in [
        "Synthesize current beauty and wellness industry context and relevant peer-reviewed marketing/customer-analytics research, while stating transfer limits.",
        "Define measurable purchase, repeat-purchase, customer-experience, and business-value outcomes before requesting data.",
        "Identify feasible internal and public data sources and define a privacy-aware variable dictionary.",
        "Formulate research questions and testable null/alternative hypotheses; keep them explicitly untested.",
        "Design a reproducible data-quality assessment and exploratory data analysis with denominators and visualization choices.",
        "Select appropriate statistical tests and conditional future segmentation/predictive methods, including assumptions, leakage safeguards, and evaluation criteria.",
        "Translate validated evidence into responsible, measurable marketing experiments and a 30–35-hour work plan.",
    ]:
        add_numbered(document, item)
    add_heading(document, "3.2 Scope and exclusions", level=2)
    add_table(
        document,
        ["In scope", "Out of scope"],
        [
            ["Approved transaction and item fields, pseudonymous linkage, category/channel, purchase cohorts, repeat behavior.", "Claims about any brand's actual performance before analysis; invented findings or survey results."],
            ["Descriptive cohort analysis and optional leakage-safe prediction if readiness criteria are met.", "Production deployment, individual automated decisions, or causal attribution from correlation."],
            ["Consent-aware campaign logs only for a separately scoped evaluation.", "Sensitive health/skin/treatment fields or inferred traits by default."],
        ],
    )
    add_body(
        document,
        "The business sponsor and data owner must settle geography, channels, "
        "brands, observation dates, test-order exclusions, order status, and "
        "financial definitions before extraction. If the historical record begins "
        "after customer acquisition, first observed purchase may not equal true "
        "first purchase.",
    )

    add_heading(document, "4. Research questions and untested hypotheses")
    add_heading(document, "4.1 Research questions", level=2)
    add_table(
        document,
        ["Question", "Variables", "Proposed analysis", "Expected output", "Business relevance"],
        [
            ["RQ1. Which product categories contribute most to repeat purchasing?", "First category; distinct qualifying orders; horizon repeat flag; cohort and complete-follow-up indicator.", "Category-level cohort rates with numerator/denominator and confidence intervals; account for cohort maturity.", "Ranked category table and interval plot, not an unsupported causal ranking.", "Identify categories warranting deeper replenishment/assortment investigation."],
            ["RQ2. How is purchase recency associated with another purchase?", "Pre-cutoff recency days; future repeat flag; tenure; category/channel covariates.", "Descriptive bins plus logistic regression as an adjusted association; nonlinear check if justified.", "Rate-by-recency plot and adjusted association with interval.", "Inform when to test a permissioned reminder; no causal claim."],
            ["RQ3. Do frequency and average order value differ across customer segments?", "Pre-cutoff segment; distinct order frequency; future-window AOV; tenure.", "Define segments before outcome window; compare with Welch ANOVA or Kruskal–Wallis/robust bootstrap as assumptions require.", "Segment summary, uncertainty and box/violin plot.", "Compare potential service needs and value without defining segments from the outcome being compared."],
            ["RQ4. Are discounts associated with order value or repeat purchase?", "Discount amount/rate; net order value; future repeat flag; category, channel, cohort.", "Descriptive cross-tabs and adjusted regression; test interactions only if pre-specified and adequately powered.", "Discount distribution and adjusted associations with intervals.", "Assess whether promotions merit controlled incremental-margin tests; discounts are confounded by selection."],
            ["RQ5. Where available, are ratings/satisfaction associated with repeat purchase?", "Linked customer rating/satisfaction; response date; future repeat flag; response status.", "Spearman/ordinal methods for ordinal ratings; logistic association only with temporally valid linked data.", "Rating-response coverage and association intervals; nonresponders described.", "Assess whether experience signals merit follow-up; do not equate satisfaction with a causal driver."],
            ["RQ6. Do purchasing patterns vary by acquisition channel or customer cohort?", "Approved acquisition channel; cohort; follow-up maturity; repeat flag/time to second order.", "Mature-cohort comparison, cross-tabs and time-to-event methods with censoring as appropriate.", "Cohort heatmap/retention curve with risk-set counts and denominators.", "Help assess channel mix and onboarding priorities, subject to attribution quality."],
            ["RQ7. Can transaction patterns support useful segmentation or future churn/repeat-risk ranking?", "Cutoff-safe RFM, category breadth, channel and future outcome at stated horizon.", "RFM rules as transparent baseline; clustering only after scaling/stability review; optional time split logistic baseline and tree comparison.", "Segment profiles/stability or temporal PR-AUC, lift-at-capacity, calibration and model card.", "Decide whether further modeling adds actionable value beyond descriptive analysis."],
        ],
        font_size=7.1,
    )
    add_heading(document, "4.2 Hypotheses (not tested)", level=2)
    hypotheses = [
        [
            ("H1 — Recency and repeat", "H0: Among customers with complete follow-up, pre-cutoff recency has no association with repeat within the chosen horizon. Ha: Shorter recency is associated with higher repeat probability."),
            ("Variables / direction", "Predictor: days since last qualifying order at cutoff. Outcome: future repeat flag. Expected direction only as a testable prior; not established."),
            ("Proposed test", "Logistic regression with pre-specified tenure/category/channel covariates; compare nonlinear form if data support it."),
            ("Assumptions / limits", "Independent customer units, reliable event dates, no leakage, complete horizon, adequate positives; association is not causation."),
        ],
        [
            ("H2 — Category differences", "H0: Horizon repeat proportions are equal across eligible first-purchase categories. Ha: At least one category proportion differs."),
            ("Variables / direction", "Predictor: first qualifying product category. Outcome: horizon repeat flag. No direction asserted."),
            ("Proposed test", "Chi-square test when expected counts/independence support it; otherwise exact/Monte Carlo alternative or combine/suppress sparse groups by prior rule."),
            ("Assumptions / limits", "Stable category definitions, mature follow-up, independent customer observations, adequate expected counts; multiple comparisons need control."),
        ],
        [
            ("H3 — Segment and future AOV", "H0: Mean future-window AOV is equal across segments defined only from pre-window behavior. Ha: At least one segment mean differs."),
            ("Variables / direction", "Group: pre-defined cutoff-safe RFM segment. Outcome: AOV in a later, fixed window. No direction asserted."),
            ("Proposed test", "Welch ANOVA with suitable post-hoc correction; Kruskal–Wallis/robust bootstrap if distribution and assumptions warrant."),
            ("Assumptions / limits", "Independent customer units, adequate follow-up; avoid circular segment definitions that include future AOV; heavy skew and unequal variance."),
        ],
        [
            ("H4 — Discount and repeat", "H0: Discount use/rate is not associated with subsequent repeat after defined adjustment. Ha: An association exists."),
            ("Variables / direction", "Predictor: pre-anchor discount amount/rate. Outcome: future horizon repeat. Direction not assumed because promotion targeting and product mix confound."),
            ("Proposed test", "Chi-square for crude eligible categorical comparison; pre-specified logistic regression for adjusted association; interaction only if powered."),
            ("Assumptions / limits", "Accurate discount basis, confounding/selection, overlap, customer-level independence; adjustment does not prove discount impact."),
        ],
        [
            ("H5 — Satisfaction and repeat", "H0: Linked satisfaction score is not associated with later repeat. Ha: A monotonic/ordered association exists."),
            ("Variables / direction", "Predictor: approved customer rating/satisfaction recorded before outcome window. Outcome: later repeat flag."),
            ("Proposed test", "Spearman rank association for ordinal exploratory analysis or ordinal/logistic regression if scale and design support it."),
            ("Assumptions / limits", "Valid instrument, correct timing/linkage, enough responses, scale consistency; nonresponse/selection bias and reverse causality."),
        ],
    ]
    for hypothesis in hypotheses:
        add_table(
            document,
            ["Hypothesis specification", "Proposed test design"],
            [[label, content] for label, content in hypothesis],
            font_size=7.8,
        )
    add_callout(
        document,
        "STATUS",
        "H1–H5 are proposed, untested hypotheses. No result is accepted or "
        "supported by customer-level evidence at this planning stage.",
    )
    add_body(
        document,
        "Before inferential testing, agree the eligible population, outcome "
        "definition, time horizon, covariates, practical effect threshold, and "
        "multiple-comparison approach with the sponsor. Exploratory comparisons "
        "must be labeled as such.",
    )
    document.add_page_break()

    add_heading(document, "5. Measures, variables, and definitions")
    add_body(
        document,
        "The analysis unit is a customer anchored to an earliest qualifying "
        "completed purchase in the available, validated history. The companion "
        "variable dictionary lists candidate source and derived fields, types, "
        "and handling caveats. Field inclusion remains subject to purpose, "
        "permissions, and data-owner review.",
    )
    add_table(
        document,
        ["Measure", "Proposed operational definition", "Important rule"],
        [
            ["Qualifying order", "Order meeting sponsor-approved status, refund, cancellation, and test-order rules.", "Apply consistently across cohorts and periods."],
            ["Repeat purchase", "At least one later qualifying order after the anchor purchase and within a stated horizon.", "Do not count cancelled/test orders; define treatment of returns."],
            ["90/180/365-day repeat rate", "Eligible anchored customers with repeat in horizon ÷ anchored customers with complete follow-up through that horizon.", "Exclude immature follow-up from denominator; report numerator and denominator."],
            ["Time to second order", "Elapsed days from first to second qualifying order.", "Censor customers without a second observed order at last valid follow-up."],
            ["Net sales", "Sales after agreed discount, refund, tax, shipping, and currency conventions.", "Finance/data-owner confirmation required; reconcile totals."],
            ["RFM features", "Recency, frequency, monetary value within a defined lookback and cutoff.", "Use fixed windows and only pre-cutoff events."],
        ],
        font_size=8.2,
    )
    add_heading(document, "5.1 Candidate variable families", level=2)
    add_table(
        document,
        ["Family", "Examples", "Use and caveat"],
        [
            ["Keys and events", "Pseudonymous customer ID, order ID, item-line ID, timestamp, status.", "Validate grain and uniqueness; no direct identifiers in the analysis layer."],
            ["Commercial", "Net sales, discount, refund, units, currency.", "Agree finance definitions; protect against invalid values and mixed currency."],
            ["Merchandise and channel", "Product/SKU, approved category, channel.", "Normalize mapping drift; include only at the granularity needed."],
            ["Derived behavior", "Cohort, order count, recency, frequency, category breadth, repeat flags.", "Document windows, cutoffs, censoring, and history completeness."],
            ["Governance / optional", "Coarse market, campaign assignment, channel-specific consent.", "Restricted; include only when approved and essential to the question."],
        ],
        font_size=8.2,
    )
    add_heading(document, "5.2 Candidate data sources", level=2)
    for item in [
        "Commerce/POS order and item-line exports for events, value, status, returns, product, and channel.",
        "CRM/customer table for approved pseudonymous linkage and minimal account fields.",
        "Product master and channel mapping for taxonomy versioning and historical consistency.",
        "Campaign/contact and consent logs only when evaluating a separately approved marketing intervention.",
        "Optional post-purchase survey or digital experience logs only if the sponsor needs customer-experience constructs and the instrument/events are approved; transaction history is not a satisfaction measure.",
        "Data governance metadata for provenance, permitted purpose, timezone, refresh cadence, and retention.",
    ]:
        add_bullet(document, item)
    add_heading(document, "5.3 Public dataset candidates and acquisition strategy", level=2)
    add_table(
        document,
        ["Candidate source", "Fields / fit", "Limitations, licensing, and access"],
        [
            [
                "UCI Machine Learning Repository, Online Retail II [11]",
                "Official page describes invoice, item code/description, quantity, invoice date, unit price, customer ID, and country. The page therefore describes genuine retail transaction-level records suitable for demonstrating purchase-history methods if permitted.",
                "Not identified as a beauty/wellness-specific retailer, so it cannot support beauty-market findings without justified product mapping. Check dataset-specific current license, attribution, access/download conditions, missing IDs, cancellation codes, and temporal representativeness before use.",
            ],
            [
                "Amazon Reviews'23 project [12]",
                "Potential product-rating, review-text, and product metadata source for permitted descriptive review analysis where a relevant category/file is confirmed.",
                "Reviews are not linked customer purchase histories and cannot establish repeat transactions. Confirm exact field schema/category availability, access process, privacy implications, and reuse/license conditions manually before downloading or processing.",
            ],
            [
                "Approved company transaction/CRM extract",
                "Best fit for retailer-specific repeat purchasing if distinct customer/order keys, dates, order outcomes, and sufficient history are available.",
                "Requires sponsor/data-owner approval, purpose/permission review, secure handling, field minimization, access controls, and an agreed retention/deletion plan.",
            ],
            [
                "Approved surveys and permitted review/social listening",
                "Could contribute satisfaction, ratings, or permitted public text signals if instruments, linkage, and purpose fit the question.",
                "Response/selection bias; non-transaction data do not prove purchasing. Do not scrape or reuse restricted platform data; check terms, consent, copyright, and applicable privacy rules.",
            ],
        ],
        font_size=7.3,
    )
    add_callout(
        document,
        "DATA AVAILABILITY",
        "No public dataset was downloaded or analyzed for this project. These are "
        "candidate sources, not evidence underlying the report. Check current "
        "licenses and access terms before any reuse.",
    )

    add_heading(document, "6. Data readiness, limitations, and privacy")
    add_heading(document, "6.1 Data limitations and quality risks", level=2)
    add_table(
        document,
        ["Risk", "Potential impact", "Planned check or treatment"],
        [
            ["Incomplete history / left truncation", "First observed order may not be true acquisition; repeat rates can be biased.", "Establish history start; qualify first-observed customers; use eligible follow-up."],
            ["Unequal follow-up / right censoring", "Recent cohorts have less time to repeat.", "Use horizon-specific complete-follow-up denominators; censor time-to-event analysis."],
            ["Identity gaps or duplicate records", "Orders split across or combined between customers.", "Validate keys, duplicates, merge rules, and linkage approval."],
            ["Returns, cancellations, test orders", "Inflated counts or misleading monetary measures.", "Document exclusions and reconcile order/refund conventions."],
            ["Taxonomy, channel, currency drift", "False subgroup differences and invalid aggregation.", "Version mapping, normalize labels, and document FX convention."],
            ["Campaign selection bias", "Recipients differ from non-recipients before contact.", "Do not claim lift from observational comparisons; use holdout/randomization."],
        ],
        font_size=8.0,
    )
    add_heading(document, "6.2 Privacy and responsible use", level=2)
    add_body(
        document,
        "Before access, confirm the approved purpose, permissions or lawful basis, "
        "minimum necessary fields, geography, retention/deletion schedule, and "
        "authorized users with the data owner and privacy/legal stakeholders. "
        "Use least-privilege access, secure transfer and storage, audit logging, "
        "pseudonymous IDs, and small-cell suppression. Purchase history is not "
        "marketing consent.",
    )
    add_body(
        document,
        "Precise location and sensitive health, skin, or treatment attributes "
        "are excluded by default. Do not infer sensitive traits from purchases. "
        "Review applicable privacy, consumer-protection, and marketing rules "
        "with qualified stakeholders; this plan is a safeguard checklist, not "
        "a legal-compliance opinion. Evaluate error rates and customer impacts "
        "across relevant groups and avoid exclusionary targeting.",
    )
    document.add_page_break()

    add_heading(document, "7. Exploratory data analysis roadmap")
    add_body(
        document,
        "The EDA starts with quality and denominators, not attractive charts. "
        "The analysis should use pandas and NumPy for validation and summaries, "
        "then Matplotlib and Seaborn for readable visual checks. No chart or "
        "statistic is generated in this planning project because no customer "
        "dataset was supplied.",
    )
    add_heading(document, "7.1 Ten-phase workflow", level=2)
    phases = [
        "Business understanding and literature review.",
        "Data-source identification, permission/licensing review, and authorized acquisition.",
        "Variable dictionary, qualifying transaction, prediction cutoff, and target definition.",
        "Data-quality assessment and source reconciliation.",
        "Documented cleaning and preprocessing with an exclusion ledger.",
        "Exploratory data analysis and cohort/segment readiness assessment.",
        "Pre-specified hypothesis testing where data assumptions support it.",
        "Interpretable RFM segmentation; optional clustering only if stable and useful.",
        "Optional repeat/churn model with time-based evaluation, customer separation, and leakage controls.",
        "Evaluation, interpretation, business economics, responsible recommendations, and limitations.",
    ]
    for index, phase in enumerate(phases, start=1):
        add_bullet(document, f"Phase {index}: {phase}")
    add_table(
        document,
        ["Stage", "Checks and analysis", "Output / decision gate"],
        [
            ["1. Governance", "Purpose, permission, minimum fields, access, retention, small-cell controls.", "Approved extract specification or pause."],
            ["2. Profile", "Schema, row grain, keys, types, missingness, date range, currency, invalid values.", "Data-quality log; documented unresolved issues."],
            ["3. Reconcile", "Order counts and sales vs approved operational totals; quantify each exclusion.", "Reconciliation table and sign-off."],
            ["4. Describe", "Order timing, distributions, spending, categories/channels; all with denominators.", "Descriptive tables/figures, coverage notes."],
            ["5. Cohorts", "Cohort sizes, mature follow-up, repeat rate by horizon, time to second order.", "Comparable cohort curves/tables; censoring caveats."],
            ["6. Compare", "Pre-specified group contrasts and intervals where assumptions allow.", "Effect sizes and uncertainty; exploratory label if not pre-specified."],
            ["7. Optional model", "Chronological split, baseline, leakage review, class and calibration metrics.", "Modeling decision memo; stop if readiness criteria fail."],
            ["8. Synthesize", "Evidence, business value, feasibility, privacy/fairness risk, test design.", "Ranked candidate experiments with owners and guardrails."],
        ],
        font_size=8.0,
    )
    add_heading(document, "7.2 Proposed visual checks", level=2)
    for item in [
        "Monthly order and net-sales time series with event/data-quality annotations.",
        "Order-frequency and monetary-value distributions, showing scale and extreme-value handling.",
        "Category and channel mix over time, with mapping changes noted.",
        "Cohort-by-follow-up-period repeat-rate matrix that suppresses immature cells.",
        "Time-to-second-order curve with at-risk counts and censoring where methodology supports it.",
        "Missingness and reconciliation summaries before any outcome interpretation.",
    ]:
        add_bullet(document, item)
    add_heading(document, "7.3 EDA question-to-output matrix", level=2)
    add_table(
        document,
        ["EDA question", "Variables / check", "Method and visualization", "Potential use"],
        [
            ["What data do we have and at what grain?", "Rows, columns, types, date range, distinct customers/orders/items.", "pandas schema/profile tables; no chart needed for initial inventory.", "Confirm feasibility and correct analytical unit."],
            ["Where are fields incomplete?", "Missing customer, date, category, price, channel; missingness patterns.", "pandas missingness counts/percentages; sorted bar chart.", "Choose eligibility and identify bias risks."],
            ["Are events duplicated or invalid?", "Duplicate transaction IDs/lines; negative/zero prices, units, dates, discount bounds.", "Rule-based counts and exception tables; histogram/boxplot for plausible numeric fields.", "Set transparent cleaning rules and quantify exclusions."],
            ["How frequent are purchases and what is the value distribution?", "Distinct orders/customer; net order value; frequency; currency.", "Summary statistics/quantiles; histogram and boxplot (log axis only if labeled).", "Understand skew, outliers, and reporting summaries."],
            ["Which categories contribute to repeat behavior?", "First category, repeat flags, cohort, maturity, category sales.", "Mature cohort pivot with denominators; bar/interval plot.", "Prioritize categories for follow-up, not assert product causality."],
            ["How do recency and frequency vary?", "Cutoff-safe recency, frequency, tenure, repeat horizon.", "RFM table and bins; recency-vs-repeat plot with confidence intervals.", "Plan future relevance or replenishment tests."],
            ["Do retention patterns vary by cohort?", "First-order cohort, elapsed follow-up, repeat event, censor date.", "Cohort retention matrix and time-to-event curve with at-risk counts.", "Avoid comparing immature cohorts; identify follow-up windows."],
            ["Do approved segments differ?", "Predefined RFM segment; future AOV/frequency/repeat.", "Box/violin and interval plots; suitable group comparison.", "Describe needs/value patterns with no causal interpretation."],
            ["Are numeric fields associated?", "Pre-specified numeric variables after quality checks.", "Pearson for suitable linear/continuous assumptions; Spearman for monotonic/ranked data; correlation heatmap only as exploratory.", "Generate testable follow-up questions; avoid causal interpretation."],
        ],
        font_size=7.4,
    )
    add_callout(
        document,
        "EDA RULE",
        "Every comparison must show its population, time window, denominator, and "
        "follow-up completeness. Visual patterns are not causal explanations.",
    )

    add_heading(document, "8. Analytical methods and evaluation")
    add_heading(document, "8.1 Descriptive and inferential methods", level=2)
    add_body(
        document,
        "Start with anchored cohort tables, horizon-based repeat rates, and "
        "time-to-event summaries. Kaplan–Meier estimates may be appropriate for "
        "time to second purchase when censoring assumptions and event definition "
        "are defensible. Group tests should match outcome type, sampling, "
        "independence, and sample size. Report effect sizes and intervals; "
        "regression adjustment can address specified covariates but does not "
        "convert an observational association into causal evidence.",
    )
    add_heading(document, "8.1.1 Method selection and assumptions", level=2)
    add_table(
        document,
        ["Method", "When it may fit", "Assumptions / caution"],
        [
            ["Descriptive statistics, frequency tables, cross-tabs", "Always begin here to understand denominators, distributions, and group composition.", "Report missingness, observation coverage, and uncertainty; exploratory patterns do not establish causes."],
            ["Chi-square test", "Association between suitable categorical variables such as first category and repeat flag.", "Independent units and adequate expected cell counts; use exact/Monte Carlo approach or justified aggregation for sparse cells."],
            ["t-test / Welch t-test; ANOVA / Welch ANOVA", "Compare means of continuous outcomes across two or several independent groups when residual/distribution behavior is acceptable.", "Independence, distribution and variance conditions; Welch handles unequal variance but not dependence or severe data problems."],
            ["Non-parametric / robust alternative", "Mann–Whitney or Kruskal–Wallis for ranked/distribution comparisons; bootstrap intervals for skewed metrics when appropriate.", "Tests may concern distributions/ranks, not means; independence and appropriate sampling still matter."],
            ["Correlation", "Explore linear Pearson or monotonic Spearman relationship between suitable numeric/ranked fields.", "Association only; outliers, nonlinear patterns, and confounding matter; do not use correlation as proof of a driver."],
            ["RFM segmentation", "Interpretable recency/frequency/monetary summaries or pre-defined bins for descriptive comparison.", "Set consistent lookbacks, adequate history, scaling/bin rules and cutoff; segments are descriptions, not inherently true customer types."],
            ["Logistic regression", "Transparent baseline for binary future-repeat outcome with interpretable adjusted associations.", "Temporal split, complete label window, calibration, collinearity/functional form, sufficient events; coefficients are not causal."],
            ["Decision tree / random forest", "Optional nonlinear benchmark if it improves a named decision over the baseline.", "Control overfitting, leakage, interpretability, calibration, stability, and temporal generalization; do not add complexity without value."],
            ["Clustering", "Optional exploratory segmentation only if volume and feature quality justify it and clusters are stable/actionable.", "Scale appropriate features, handle skew/missingness, assess stability and usefulness; silhouette alone does not establish business value."],
        ],
        font_size=7.4,
    )
    add_heading(document, "8.2 Optional predictive modeling", level=2)
    add_body(
        document,
        "Model only if a named business decision needs ranking or probability "
        "estimates and data-readiness criteria pass. Define the prediction "
        "cutoff, eligible customer population, future outcome horizon, and "
        "positive class first. Use a chronological train/validation/test split, "
        "separate customers across partitions, and exclude every field created "
        "after the cutoff. Fit preprocessing inside training folds. Compare "
        "scikit-learn logistic regression to a prevalence-only baseline before "
        "considering a more complex estimator.",
    )
    add_table(
        document,
        ["Metric", "Proposed use", "Decision caveat"],
        [
            ["Repeat rate + denominator", "Communicate cohort outcomes at a named horizon.", "Only compare populations with complete, comparable follow-up."],
            ["Accuracy / precision / recall / F1", "Report threshold-dependent classification behavior where meaningful.", "Accuracy alone can mislead under class imbalance; choose threshold based on action costs."],
            ["Precision-recall AUC", "Summarize ranking under class imbalance.", "Does not pick a business threshold or show calibration."],
            ["ROC AUC", "Secondary measure of ranking discrimination.", "Can obscure positive-class performance when prevalence is low."],
            ["Precision, recall, lift at capacity", "Evaluate quality among the number of customers the team can act on.", "Capacity, costs, and customer safeguards define useful thresholds."],
            ["Log loss / calibration / Brier score", "Evaluate probability quality if scores will be interpreted as probabilities.", "Calibration must be checked on temporally held-out data."],
            ["Incremental margin + guardrails", "Evaluate intervention economics and experience in an experiment.", "Use randomized holdout where feasible; define opt-out/complaint safeguards."],
        ],
        font_size=8.0,
    )
    add_body(
        document,
        "No universal pass threshold is assumed. Agree business capacity, cost, "
        "minimum practical effect, uncertainty, sample-size method, and stop "
        "conditions before seeing holdout results. The customer-analytics "
        "literature [5]–[7] motivates these candidate methods, but findings and "
        "model assumptions must be tested locally. Tool documentation for "
        "planned evaluation metrics is listed in reference [8].",
    )
    add_heading(document, "8.3 Expected value and proposed business KPIs", level=2)
    add_table(
        document,
        ["KPI family", "Candidate measure", "Interpretation safeguard"],
        [
            ["Customer retention", "90/180/365-day repeat rate; time to second qualifying purchase.", "Predefine horizon and complete-follow-up denominator; no benchmark is presumed."],
            ["Purchase/value", "Orders per eligible customer; net AOV; contribution margin after promotion/contact cost.", "Agree finance definitions, returns, currency and tax treatment."],
            ["Marketing effectiveness", "Randomized incremental repeat rate and incremental contribution margin.", "Use holdout when feasible; observational association is not campaign lift."],
            ["Customer guardrails", "Opt-out, complaint, return, unsubscribe, and approved experience measures.", "Set stop conditions before campaign; use only permitted contact channels."],
            ["Predictive usefulness", "Lift/precision at contact capacity, calibration, and temporal stability.", "Score quality alone does not show intervention value or permission to contact."],
        ],
        font_size=7.6,
    )
    document.add_page_break()

    add_heading(document, "9. Business recommendation and test framework")
    add_body(
        document,
        "Recommendations will be considered only after data validation and "
        "interpretation. Score each candidate on evidence strength, customer "
        "benefit, incremental business value, operational feasibility, "
        "privacy/fairness risk, and testability. Keep the evidence chain visible: "
        "observation → interpretation → decision hypothesis → test → scale/stop.",
    )
    add_table(
        document,
        ["Experiment-card field", "Required specification"],
        [
            ["Decision and hypothesis", "What action is being tested and why; distinguish evidence from assumption."],
            ["Population and exclusions", "Eligibility, consent, channel reachability, and exclusion/suppression rules."],
            ["Treatment and control", "Intervention, randomization unit, holdout, exposure logging, and contamination risks."],
            ["Outcome and guardrails", "Primary incremental outcome plus margin, opt-out, complaint, and customer-experience safeguards."],
            ["Analysis and duration", "Sample-size/power method, observation window, predeclared analysis, and stopping rule."],
            ["Ownership", "Decision owner, execution owner, review date, and criteria to expand, revise, or stop."],
        ],
    )
    add_heading(document, "9.1 Candidate interventions to evaluate", level=2)
    add_body(
        document,
        "Potential test concepts—not recommendations based on observed customer "
        "results—include helpful onboarding, category-relevant product "
        "education, permissioned replenishment reminders, or a service follow-up. "
        "Avoid blanket discounts as the default: evaluate incremental contribution "
        "margin and customer experience against a holdout. Do not contact people "
        "without valid channel permission.",
    )
    add_callout(
        document,
        "CAUSAL CLAIMS",
        "A predictive score identifies a pattern; it does not show that contacting "
        "a scored customer will cause a purchase. Use a well-designed experiment "
        "to estimate intervention impact.",
    )

    add_heading(document, "10. Tools, timeline, and project governance")
    add_body(
        document,
        "The proposed toolchain is Python in VS Code (or Jupyter), pandas and "
        "NumPy for data preparation and measurement, Matplotlib and Seaborn for "
        "EDA, scikit-learn for conditional future predictive modeling, and "
        "python-docx for this report. The report generator itself reads no "
        "customer data. Official tool references are logged in [8]–[10].",
    )
    add_table(
        document,
        ["Workstream", "Hours", "Planned output / gate"],
        [
            ["Business framing and source review", "4", "Decision question, source log, scope."],
            ["Governance and data inventory", "4", "Approved field list, access and readiness checklist."],
            ["Data extraction and quality checks", "6", "Reconciled, documented analytical table."],
            ["EDA and cohort definitions", "7", "Summary tables, charts, cohort measures."],
            ["Optional exploratory modeling", "5", "Baseline, chronological split, metrics only if gate passes."],
            ["Interpretation and test design", "4", "Evidence-to-action experiment cards."],
            ["Report, QA, and presentation", "4", "Reviewed report and stakeholder handoff."],
            ["TOTAL", "34", "Within 30–35-hour target."],
        ],
        font_size=8.2,
    )
    add_body(
        document,
        "The five-hour modeling allocation is conditional. If permissions, "
        "follow-up, sample size, or data quality are insufficient, redirect that "
        "time to data remediation, uncertainty analysis, or stakeholder "
        "validation. The project sponsor approves scope; the data owner approves "
        "definitions/access; a privacy or legal reviewer advises on applicable "
        "requirements; the analyst documents code, decisions, and limitations.",
    )
    document.add_page_break()

    add_heading(document, "11. Expected deliverables and success criteria")
    add_table(
        document,
        ["Deliverable", "Acceptance check"],
        [
            ["Approved analysis specification", "Population, horizon, qualifying order, metric definitions, and exclusions documented."],
            ["Validated analytical table and QA log", "Grain, joins, totals, exclusions, timezones, missingness, and currencies reviewed."],
            ["EDA and cohort outputs", "Reproducible code, visible denominators, maturity/censoring notes, and restrained interpretation."],
            ["Optional model assessment", "Readiness gate passed; temporal holdout, leakage check, baseline, calibration/ranking, and intended use documented."],
            ["Test-ready recommendation cards", "Action, control, primary outcome, guardrails, owners, and stop/scale criteria specified."],
            ["Stakeholder report and handoff", "Sources, assumptions, caveats, decisions, and next actions reviewed."],
        ],
    )
    add_heading(document, "11.1 Week 1 exit criteria", level=2)
    for item in [
        "Sponsor agrees the decision question, population, horizon, and success criteria.",
        "Data owner approves the source, minimum field set, and order/revenue definitions.",
        "Privacy/access review is complete before any customer extract is used.",
        "Follow-up completeness and cohort eligibility can be determined.",
        "Next phase has an accountable owner and a stop/go decision for optional modeling.",
    ]:
        add_bullet(document, item)

    add_heading(document, "12. Conclusion")
    add_body(
        document,
        "This Week 1 project establishes a defensible route from industry "
        "questions to customer-level evidence: approve the purpose and data, "
        "define purchase and repeat consistently, validate and reconcile the "
        "history, compare cohorts with fair follow-up, and test any proposed "
        "action against a holdout. Beauty and wellness market publications "
        "and academic customer-analytics research provide context and "
        "methodological grounding, not a substitute for the business's own "
        "validated evidence.",
    )
    add_body(
        document,
        "At this stage, every customer-level question and hypothesis remains "
        "open. There are no observed customer findings, no model, and no proven "
        "retention intervention. The appropriate immediate next step is "
        "stakeholder approval of the decision question and governance gate, "
        "followed by a documented data inventory.",
    )
    document.add_page_break()

    add_heading(document, "Appendix A. Detailed variable dictionary")
    add_body(
        document,
        "The accompanying CSV is the authoritative machine-readable dictionary. "
        "This appendix reproduces its definitions, examples, analytical roles, "
        "sources, inclusion status, quality checks, and privacy notes for review "
        "within the report. All fields are proposed; availability and permitted "
        "use have not been confirmed.",
    )
    dictionary = load_variable_dictionary()
    dictionary_rows = []
    for item in dictionary:
        dictionary_rows.append(
            [
                item["variable"],
                item["definition"],
                item["data_type_example"],
                f'{item["role"]}; source: {item["source"]}; {item["required_optional"]}',
                f'Quality: {item["quality_checks"]} Privacy: {item["privacy_notes"]}',
            ]
        )
    add_table(
        document,
        ["Variable", "Definition", "Type / example", "Role / source / status", "Quality / privacy"],
        dictionary_rows,
        font_size=6.1,
    )
    document.add_page_break()

    add_heading(document, "References")
    references = [
        (
            "[1] McKinsey & Company. (2026, June 18). From aisle to algorithm: "
            "The beauty categories, channels, and concepts shaping 2030 growth.",
            "https://www.mckinsey.com/industries/consumer-packaged-goods/our-insights/state-of-beauty",
        ),
        (
            "[2] McKinsey & Company. (2025, June 9). State of Beauty 2025: "
            "Solving a shifting growth puzzle.",
            "https://www.mckinsey.com/industries/consumer-packaged-goods/our-insights/state-of-beauty-2025",
        ),
        (
            "[3] McKinsey & Company. (2025, May 29). The $2 trillion global "
            "wellness market gets a millennial and Gen Z glow-up.",
            "https://www.mckinsey.com/industries/consumer-packaged-goods/our-insights/future-of-wellness-trends",
        ),
        (
            "[4] Perret, J. K., & Schwientek, J. (2025). Beauty Tech—Customer "
            "Experience and Loyalty of Augmented Reality- and Artificial "
            "Intelligence-Driven Cosmetics. Digital, 5(2), 21.",
            "https://doi.org/10.3390/digital5020021",
        ),
        (
            "[5] Fader, P. S., Hardie, B. G. S., & Lee, K. L. (2005). RFM and "
            "CLV: Using Iso-Value Curves for Customer Base Analysis. Journal of "
            "Marketing Research, 42(4), 415–430.",
            "https://doi.org/10.1509/jmkr.2005.42.4.415",
        ),
        (
            "[6] Fader, P. S., Hardie, B. G. S., & Shang, J. (2010). "
            "Customer-Base Analysis in a Discrete-Time Noncontractual Setting. "
            "Marketing Science, 29(6), 1086–1108.",
            "https://doi.org/10.1287/mksc.1100.0580",
        ),
        (
            "[7] Gupta, S., Lehmann, D. R., & Stuart, J. A. (2004). Valuing "
            "Customers. Journal of Marketing Research, 41(1), 7–18.",
            "https://doi.org/10.1509/jmkr.41.1.7.25084",
        ),
        (
            "[8] scikit-learn developers. Model evaluation: quantifying the "
            "quality of predictions. Official documentation.",
            "https://scikit-learn.org/stable/modules/model_evaluation.html",
        ),
        (
            "[9] pandas development team. pandas documentation.",
            "https://pandas.pydata.org/docs/",
        ),
        (
            "[10] python-docx contributors. python-docx documentation.",
            "https://python-docx.readthedocs.io/en/latest/",
        ),
        (
            "[11] UCI Machine Learning Repository. Online Retail II. Dataset "
            "record and variable descriptions; publication date not stated on "
            "the accessed page. Dataset-specific reuse terms require review.",
            "https://archive.ics.uci.edu/dataset/502/online+retail+ii",
        ),
        (
            "[12] Amazon Reviews'23 project. Dataset documentation; publication "
            "date not stated on the accessed landing page. Confirm category, "
            "fields, access, and reuse terms before use.",
            "https://amazon-reviews-2023.github.io/",
        ),
    ]
    for citation, url in references:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run(citation + " ")
        add_hyperlink(paragraph, url, url)
    add_body(
        document,
        "Access dates: McKinsey public pages October 10, 2026; academic DOI "
        "records and abstracts October 11, 2026. Bibliographic details for "
        "references [4]–[7] were checked against Crossref DOI records. The "
        "publisher page for [4] returned HTTP 403 during this review, so its "
        "use here is limited to the abstract-level information available from "
        "the DOI record; consult the full article for study-design details. "
        "Industry forecasts and survey results remain attributed to their "
        "publishers and are not recreated or independently estimated here.",
    )

    add_heading(document, "Final quality check")
    checks = [
        ("Source dates and titles", "McKinsey dates verified on public pages; academic citation metadata checked against DOI records."),
        ("Industry and academic evidence", "Beauty-specific academic context and customer-analytics methods are discussed with transfer limitations."),
        ("In-text citations", "Industry statements cite [1]–[3]; academic concepts cite [4]–[7]; tools cite [8]–[10]; candidate data sources cite [11]–[12]."),
        ("Customer findings", "None reported; no customer dataset was provided or analyzed."),
        ("Hypotheses", "H1–H5 each include null/alternative framing, variables, candidate test, and limitations; all are untested."),
        ("Methods and metrics", "Population, follow-up, censoring, leakage, and decision metrics are specified as proposals."),
        ("Privacy", "Data minimization, permissions, restricted fields, and reviewer gate are stated; no compliance guarantee is made."),
        ("Deliverable structure", "Cover, TOC field, numbered headings, tables, references, header/footer, and page fields included."),
        ("Companion files", "Research source log, variable dictionary, analysis plan, README, and requirements are part of the project."),
    ]
    add_table(document, ["Check", "Result"], [list(row) for row in checks], font_size=8.5)
    add_callout(
        document,
        "FINAL STATUS",
        "Proposal complete. Customer analysis and all hypotheses remain untested.",
    )


def build_report() -> Path:
    global TABLE_COUNTER
    TABLE_COUNTER = 0
    validate_source_log()
    document = Document()
    set_document_styles(document)
    configure_section(document.sections[0])
    set_update_fields(document)
    document.core_properties.title = (
        "Beauty & Wellness Customer Purchase and Retention Analysis — Week 1"
    )
    document.core_properties.subject = (
        "Proposed customer purchase and retention analysis; no findings tested"
    )
    document.core_properties.author = "Internship Project"
    document.core_properties.keywords = (
        "beauty, wellness, customer retention, cohort analysis, proposal"
    )

    add_cover(document)
    add_table_of_contents(document)
    add_main_content(document)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    verify_docx(OUTPUT_PATH)
    return OUTPUT_PATH


def verify_docx(path: Path) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"Generated report is missing or empty: {path}")
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        required_parts = {"word/document.xml", "word/styles.xml"}
        if not required_parts.issubset(names):
            raise RuntimeError("Generated file is missing required Word document parts.")
        content = archive.read("word/document.xml").decode("utf-8")
        relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        required_text = (
            "Executive summary",
            "Hypotheses (not tested)",
            "Academic research and methodological rationale",
            "Public dataset candidates and acquisition strategy",
            "Detailed variable dictionary",
            "Final quality check",
            "References",
        )
        missing = [text for text in required_text if text not in content]
        if missing:
            raise RuntimeError(f"Generated report is missing sections: {missing}")
        if "TOC" not in content:
            raise RuntimeError("The generated report is missing its table-of-contents field.")
        expected_urls = (
            "https://www.mckinsey.com/industries/consumer-packaged-goods/our-insights/state-of-beauty",
            "https://www.mckinsey.com/industries/consumer-packaged-goods/our-insights/state-of-beauty-2025",
            "https://www.mckinsey.com/industries/consumer-packaged-goods/our-insights/future-of-wellness-trends",
            "https://doi.org/10.3390/digital5020021",
            "https://doi.org/10.1509/jmkr.2005.42.4.415",
            "https://doi.org/10.1287/mksc.1100.0580",
            "https://doi.org/10.1509/jmkr.41.1.7.25084",
            "https://archive.ics.uci.edu/dataset/502/online+retail+ii",
            "https://amazon-reviews-2023.github.io/",
        )
        missing_urls = [url for url in expected_urls if url not in relationships]
        if missing_urls:
            raise RuntimeError(f"Generated report is missing linked references: {missing_urls}")
        footer_parts = [name for name in names if name.startswith("word/footer") and name.endswith(".xml")]
        if not footer_parts or not any(
            "PAGE" in archive.read(name).decode("utf-8") for name in footer_parts
        ):
            raise RuntimeError("The generated report is missing its page-number field.")
    load_variable_dictionary()


if __name__ == "__main__":
    try:
        report_path = build_report()
    except (FileNotFoundError, ValueError, RuntimeError) as error:
        raise SystemExit(f"Report generation failed: {error}") from error
    print(f"Generated and verified: {report_path}")
    print(f"File size: {report_path.stat().st_size:,} bytes")
